#!/usr/bin/env python3
"""Create a reproducible inventory of the analysis workspace.

The inventory is intentionally read-only with respect to user data.  Raw crawl
trees are summarized by count and byte size because their own manifests already
carry per-resource SHA-256 hashes; other files are hashed individually so exact
duplicates and obsolete generated artifacts can be reviewed before archiving.
"""

from __future__ import annotations

import ast
import csv
import hashlib
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


WORKSPACE = Path(__file__).resolve().parents[1]
METADATA = WORKSPACE / "metadata"
JSON_OUT = METADATA / "workspace_inventory.json"
CSV_OUT = METADATA / "workspace_inventory_files.csv"

SUMMARY_ONLY_DIRS = {
    "data_raw",
    ".venv",
    "cache",
    "cache_data",
    "archive_legacy",
    "staging",
}
EXCLUDE_DIRS = {
    ".git",
    ".review_media",
    ".tmp_docx_audit_v352",
    "__pycache__",
    # 本地构建依赖和发布中间目录不纳入工作区清单；它们可由构建脚本重新生成。
    ".build_python",
    ".build_tcl",
    ".build_tk",
    ".pyinstaller_vendor",
    "build",
    "build_single",
    "dist",
}
EXCLUDE_SUFFIXES = {".lock"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    return path.relative_to(WORKSPACE).as_posix()


def is_excluded_file(path: Path) -> bool:
    """判断文件是否位于不应进入清单的缓存/构建目录。"""
    parts = set(path.relative_to(WORKSPACE).parts)
    return bool(parts & EXCLUDE_DIRS) or path.suffix.lower() in EXCLUDE_SUFFIXES


def iter_nonraw_files():
    for path in WORKSPACE.rglob("*"):
        if not path.is_file():
            continue
        if is_excluded_file(path) or path.relative_to(WORKSPACE).parts[0] in SUMMARY_ONLY_DIRS:
            continue
        if path in {JSON_OUT, CSV_OUT}:
            continue
        yield path


def workbook_summary(path: Path) -> dict:
    item = {"path": rel(path), "status": "ok", "sheets": []}
    try:
        book = load_workbook(path, read_only=True, data_only=False)
        for sheet in book.worksheets:
            header = []
            first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            for value in first:
                header.append(None if value is None else str(value))
            item["sheets"].append(
                {
                    "name": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "header": header,
                }
            )
        book.close()
    except Exception as exc:  # damaged/unsupported workbooks remain visible
        item["status"] = "error"
        item["error"] = f"{type(exc).__name__}: {exc}"
    return item


def python_summary(path: Path) -> dict:
    result = {"path": rel(path), "syntax_ok": True, "imports": []}
    try:
        tree = ast.parse(path.read_text(encoding="utf-8-sig"), filename=rel(path))
        imports = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imports.update(alias.name.split(".")[0] for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imports.add(node.module.split(".")[0])
        result["imports"] = sorted(imports)
    except Exception as exc:
        result["syntax_ok"] = False
        result["error"] = f"{type(exc).__name__}: {exc}"
    return result


def main() -> int:
    files = []
    duplicate_index: dict[tuple[int, str], list[str]] = defaultdict(list)
    for path in sorted(iter_nonraw_files(), key=lambda value: rel(value).casefold()):
        stat = path.stat()
        digest = sha256(path)
        record = {
            "path": rel(path),
            "bytes": stat.st_size,
            "modified_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
            "sha256": digest,
            "suffix": path.suffix.lower(),
            "top_directory": path.relative_to(WORKSPACE).parts[0],
        }
        files.append(record)
        duplicate_index[(stat.st_size, digest)].append(record["path"])

    directory_summaries = []
    for child in sorted(WORKSPACE.iterdir(), key=lambda value: value.name.casefold()):
        if not child.is_dir() or child.name in EXCLUDE_DIRS:
            continue
        members = [path for path in child.rglob("*") if path.is_file() and not is_excluded_file(path)]
        directory_summaries.append(
            {"path": child.name, "file_count": len(members), "bytes": sum(path.stat().st_size for path in members)}
        )

    duplicate_groups = [
        {"bytes": size, "sha256": digest, "paths": paths}
        for (size, digest), paths in duplicate_index.items()
        if len(paths) > 1
    ]
    workbooks = [workbook_summary(WORKSPACE / item["path"]) for item in files if item["suffix"] == ".xlsx"]
    python_files = [python_summary(WORKSPACE / item["path"]) for item in files if item["suffix"] == ".py"]

    payload = {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "workspace": ".",
        "raw_note": "data_raw is summarized by directory; per-resource hashes are in crawl manifests",
        "excluded_note": "Transient runtime .lock files and excluded/cache directories are not hashed",
        "files": files,
        "directories": directory_summaries,
        "exact_duplicate_groups": sorted(duplicate_groups, key=lambda item: (-len(item["paths"]), item["paths"][0])),
        "workbooks": workbooks,
        "python_files": python_files,
    }
    METADATA.mkdir(parents=True, exist_ok=True)
    JSON_OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with CSV_OUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["path", "bytes", "modified_at", "sha256", "suffix", "top_directory"])
        writer.writeheader()
        writer.writerows(files)
    print(json.dumps({"files": len(files), "workbooks": len(workbooks), "duplicate_groups": len(duplicate_groups)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
